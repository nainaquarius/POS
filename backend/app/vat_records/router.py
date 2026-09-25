from io import BytesIO

from fastapi import APIRouter, Depends, Query
from fastapi.responses import StreamingResponse

from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, Border, Side

from app.dependencies import require_roles
from app.vat_records.schema import VatRecordResponse
from app.vat_records.service import getVatRecords


router = APIRouter(prefix="/vat", tags=["Vat"])


@router.get("/", response_model=list[VatRecordResponse])
async def get_vat_records(
    from_date: str | None = Query(default=None),
    to_date: str | None = Query(default=None),
    current_user=Depends(
        require_roles("admin", "manager", "cashier")
    )
):
    return await getVatRecords(from_date, to_date)


@router.get("/export")
async def export_vat_records(
    from_date: str | None = Query(default=None),
    to_date: str | None = Query(default=None),
    current_user=Depends(
        require_roles("admin", "manager", "cashier")
    )
):
    records = await getVatRecords(from_date, to_date)

    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = "VAT Records"

    # Headers
    headers = [
        "S.No.",
        "City Name",
        "Products",
        "Invoice Number",
        "Date",
        "Amount Without VAT",
        "VAT %",
        "Total Amount",
        "VAT Amount",
    ]

    worksheet.append(headers)

    # Header styling
    for cell in worksheet[1]:
        cell.font = Font(bold=True)
        cell.alignment = Alignment(
            horizontal="center",
            vertical="center"
        )

    # Add records
    for index, record in enumerate(records, start=1):

        worksheet.append([
            index,
            "Ajman",
            record.product_names,
            record.invoice_number,
            record.date.strftime("%d-%m-%Y"),
            record.invoice_price,
            0.05,
            record.total_amount,
            record.vat,
        ])

    # Total row
    total_row = worksheet.max_row + 1

    worksheet.cell(row=total_row, column=5, value="TOTAL")

    worksheet.cell(
        row=total_row,
        column=6,
        value=f"=SUM(F2:F{total_row - 1})"
    )

    worksheet.cell(
        row=total_row,
        column=8,
        value=f"=SUM(H2:H{total_row - 1})"
    )

    worksheet.cell(
        row=total_row,
        column=9,
        value=f"=SUM(I2:I{total_row - 1})"
    )

    # Total row styling
    for cell in worksheet[total_row]:
        cell.font = Font(bold=True)

    # Number formatting
    for row in range(2, total_row):
        worksheet.cell(row=row, column=6).number_format = "0.00"
        worksheet.cell(row=row, column=7).number_format = "0%"
        worksheet.cell(row=row, column=8).number_format = "0.00"
        worksheet.cell(row=row, column=9).number_format = "0.00"

    worksheet.cell(
        row=total_row,
        column=6
    ).number_format = "0.00"

    worksheet.cell(
        row=total_row,
        column=8
    ).number_format = "0.00"

    worksheet.cell(
        row=total_row,
        column=9
    ).number_format = "0.00"

    # Borders
    thin_border = Border(
        left=Side(style="thin"),
        right=Side(style="thin"),
        top=Side(style="thin"),
        bottom=Side(style="thin"),
    )

    for row in worksheet.iter_rows():
        for cell in row:
            cell.border = thin_border
            cell.alignment = Alignment(
                vertical="center"
            )

    # Center selected columns
    for row in range(1, total_row + 1):
        worksheet.cell(row=row, column=1).alignment = Alignment(
            horizontal="center"
        )

        worksheet.cell(row=row, column=2).alignment = Alignment(
            horizontal="center"
        )

        worksheet.cell(row=row, column=4).alignment = Alignment(
            horizontal="center"
        )

        worksheet.cell(row=row, column=5).alignment = Alignment(
            horizontal="center"
        )

        worksheet.cell(row=row, column=7).alignment = Alignment(
            horizontal="center"
        )

    # Column widths
    worksheet.column_dimensions["A"].width = 8
    worksheet.column_dimensions["B"].width = 15
    worksheet.column_dimensions["C"].width = 30
    worksheet.column_dimensions["D"].width = 18
    worksheet.column_dimensions["E"].width = 15
    worksheet.column_dimensions["F"].width = 22
    worksheet.column_dimensions["G"].width = 10
    worksheet.column_dimensions["H"].width = 18
    worksheet.column_dimensions["I"].width = 18


    worksheet.freeze_panes = "A2"

    output = BytesIO()
    workbook.save(output)
    output.seek(0)

    return StreamingResponse(
        output,
        media_type=(
            "application/vnd.openxmlformats-officedocument"
            ".spreadsheetml.sheet"
        ),
        headers={
            "Content-Disposition": (
                'attachment; filename="vat_records.xlsx"'
            )
        }
    )